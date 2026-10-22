"""마켓 카테고리 목록 갱신: 이셀러스 [마켓카테고리 엑셀다운] 파일(OpenMarketCategory.xlsx)로 market_categories.json.gz를 새로 만듭니다.
사용: python build_market_db.py <OpenMarketCategory.xlsx>
받는 곳: 이셀러스 원본상품 > 원본상품 엑셀추가 > [마켓카테고리 엑셀다운]
'카테고리코드'로 시작하는 마켓 시트(옥션, G마켓, 11번가, 스마트스토어, 쿠팡, 롯데ON, 톡스토어 등)만 읽습니다."""
import gzip
import json
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import MARKET_CAT_PATH  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

book = load_workbook(sys.argv[1], read_only=True, data_only=True)
markets, updated = {}, {}
for sheet in book.worksheets:
    rows = sheet.iter_rows(values_only=True)
    head = [str(h or "").replace(" ", "") for h in next(rows, ())]
    if sheet.title == "정보":
        for row in rows:
            if row and len(row) > 2 and row[1] and row[2] and row[1] != "마켓":
                updated[str(row[1])] = str(row[2])[:10]
        continue
    if not head or head[0] != "카테고리코드" or any(w in sheet.title for w in ("옵션", "속성")):
        continue  # 추천옵션·구매옵션·상품속성 시트, ESM2.0처럼 형식이 다른 시트는 건너뜀
    names = [i for i, h in enumerate(head) if h.endswith("카테고리명")][:6]
    cats = {}
    for row in rows:
        if row and row[0]:
            cats[str(row[0]).strip()] = " > ".join(str(row[i]).strip() for i in names if row[i])
    markets[sheet.title] = cats
if "스마트스토어" not in markets:
    sys.exit("마켓 시트를 찾지 못했습니다. 이셀러스 마켓카테고리 엑셀이 맞는지 확인하세요.")
if MARKET_CAT_PATH.exists():
    shutil.copy(MARKET_CAT_PATH, MARKET_CAT_PATH.with_name("market_categories.backup.json.gz"))
with gzip.open(MARKET_CAT_PATH, "wt", encoding="utf-8") as h:
    json.dump({"updated": updated, "markets": markets}, h, ensure_ascii=False)
print("마켓 카테고리 목록을 새로 만들었습니다:", ", ".join(f"{k} {len(v):,}개" for k, v in markets.items()))
