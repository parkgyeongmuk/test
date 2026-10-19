"""키워드 DB 갱신: 아이템스카우트에서 받은 엑셀(아이템발굴/키워드분석 다운로드 파일, 여러 개 가능)로 keyword_db.csv.gz를 새로 만듭니다.
사용: python build_db.py <엑셀 또는 폴더> [<엑셀 또는 폴더> ...]
같은 키워드는 총 검색수가 가장 큰 행을 남깁니다. 기존 DB는 keyword_db.backup.csv.gz로 보관합니다."""
import csv
import gzip
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import DB_PATH, compact  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def num(v):
    try:
        return float(str(v).replace(",", "")) if v not in (None, "", "-") else None
    except ValueError:
        return None


files = []
for arg in sys.argv[1:]:
    p = Path(arg)
    files += sorted(p.glob("*.xlsx")) if p.is_dir() else [p]
best = {}
for f in files:
    sheet = load_workbook(f, read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    head = None
    for _ in range(5):
        row = next(rows, None)
        if row and "키워드" in row and any(h in row for h in ("총 검색수", "검색수")):
            head = [str(h or "").strip() for h in row]
            break
    if not head:
        print("건너뜀(형식 다름):", f.name)
        continue
    col = {h: i for i, h in enumerate(head) if h}
    vol_col = col.get("총 검색수", col.get("검색수"))
    for row in rows:
        if not row or not row[col["키워드"]]:
            continue
        kw = str(row[col["키워드"]]).strip()
        key = compact(kw)
        vol = num(row[vol_col]) or 0
        if not key or (key in best and best[key][2] >= vol):
            continue
        prod = num(row[col["상품수"]]) if "상품수" in col else None
        comp = num(row[col["경쟁강도"]]) if "경쟁강도" in col else None
        cat = str(row[col["대표 카테고리"]] or "").strip() if "대표 카테고리" in col else ""
        best[key] = [kw, cat, int(vol), "" if prod is None else int(prod), "" if comp is None else round(comp, 2)]
if not best:
    sys.exit("키워드를 하나도 읽지 못했습니다. 아이템스카우트 엑셀이 맞는지 확인하세요.")
if DB_PATH.exists():
    shutil.copy(DB_PATH, DB_PATH.with_name("keyword_db.backup.csv.gz"))
with gzip.open(DB_PATH, "wt", encoding="utf-8", newline="") as h:
    w = csv.writer(h)
    w.writerow(["keyword", "category", "volume", "products", "competition"])
    for row in sorted(best.values(), key=lambda r: -r[2]):
        w.writerow(row)
print(f"키워드 {len(best):,}개로 DB를 새로 만들었습니다 ({len(files)}개 파일).")
