"""Shared helpers: eSellers Excel I/O, Itemscout keyword DB, title rules."""
from __future__ import annotations

import csv
import gzip
import html
import io
import json
import math
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parent.parent
DB_PATH = SKILL_DIR / "keyword_db.csv.gz"
CONFIG_PATH = SKILL_DIR / "config.json"
MARKET_CAT_PATH = SKILL_DIR / "market_categories.json.gz"
TITLE_MIN, TITLE_MAX = 30, 40

COLOR_WORDS = ("블랙", "화이트", "아이보리", "베이지", "레드", "블루", "그린", "핑크", "브라운", "그레이", "네이비",
               "실버", "골드", "투명", "검정", "흰색", "회색", "빨강", "파랑", "노랑", "옐로우", "퍼플", "보라", "민트", "카키")
SIZE_WORDS = ("대형", "중형", "소형", "미니", "라지", "스몰", "대용량", "소용량", "특대", "초대형")
PROMO_WORDS = ("최고", "특가", "인기", "추천", "간편", "핫딜", "초특가", "best", "베스트", "신상", "최저가")


def config() -> dict:
    try:
        return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


DEFAULT_MARKETS = [{"sheet": "스마트스토어", "name": "스마트스토어"}, {"sheet": "11번가", "name": "11번가"},
                   {"sheet": "옥션", "name": "옥션"}, {"sheet": "G마켓", "name": "지마켓"},
                   {"sheet": "톡스토어", "name": "톡스토어"}]


def markets() -> list[dict]:
    """확장정보 '마켓 카테고리번호'에 적을 마켓. sheet = 마켓카테고리 엑셀 시트 이름, name = 이셀러스에 적는 마켓명."""
    return config().get("markets", DEFAULT_MARKETS)


def brands() -> list[str]:
    return config().get("brands", ["사이노텍", "다이소", "스탠리", "무인양품", "이케아", "락앤락", "코스트코", "스타벅스"])


def compact(value) -> str:
    return re.sub(r"[^가-힣A-Za-z0-9]", "", html.unescape(str(value or ""))).lower()


def clean(value, brand_list: list[str] | None = None) -> str:
    text = html.unescape(str(value or "")).replace("억세사리", "액세서리")
    for brand in brand_list if brand_list is not None else brands():
        text = re.sub(re.escape(brand), " ", text, flags=re.I)
    text = re.sub(r"[^가-힣A-Za-z0-9\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


_market_cats = None


def market_categories() -> dict[str, dict[str, str]]:
    """마켓(시트 이름) → {카테고리 번호: 경로}. 이셀러스 마켓카테고리 엑셀로 만든 목록입니다."""
    global _market_cats
    if _market_cats is None:
        _market_cats = {}
        if MARKET_CAT_PATH.exists():
            with gzip.open(MARKET_CAT_PATH, "rt", encoding="utf-8") as handle:
                _market_cats = json.load(handle)["markets"]
    return _market_cats


def naver_categories() -> dict[str, str]:
    """스마트스토어(네이버쇼핑) 카테고리 번호 → 경로('생활/건강 > 주방용품 > ...')."""
    return market_categories().get("스마트스토어", {})


def naver_code(path: str, cats: dict[str, str]) -> str:
    """경로가 목록에 정확히 있으면 번호. 없으면 마지막 이름이 같은 카테고리가 하나뿐일 때 그 번호, 아니면 ''."""
    key = compact(path)
    exact = next((code for code, p in cats.items() if compact(p) == key), "")
    if exact:
        return exact
    leaf = compact(path.split(">")[-1])
    same = [code for code, p in cats.items() if compact(p.split(">")[-1]) == leaf]
    return same[0] if len(same) == 1 else ""


# ---------------------------------------------------------------- Excel (기본정보 시트)

def _norm(h) -> str:
    return re.sub(r"[\s*()（）/\\_\-·]+", "", str(h)).lower()


def find_columns(headers: list, offset: int = 0) -> dict | None:
    normalized = [_norm(h) for h in headers]

    def find(pred):
        return next((i + offset for i, h in enumerate(normalized) if pred(h)), None)

    cols = {"name": find(lambda h: h.startswith("상품명")), "price": find(lambda h: h.startswith("판매가")),
            "quantity": find(lambda h: h.startswith("수량")), "origin": find(lambda h: h.startswith("원산지")),
            "folder": find(lambda h: h.startswith("폴더명")), "no": find(lambda h: h.startswith("원본번호")),
            "option_details": find(lambda h: "선택사항" in h and "상세정보" in h)}
    return cols if cols["name"] is not None else None


def read_rows(path: Path) -> tuple[list[dict], dict]:
    data = path.read_bytes()
    if path.suffix.lower() == ".xls":
        import xlrd
        sheet = xlrd.open_workbook(file_contents=data).sheet_by_name("기본정보")
        get = lambda r, c: sheet.cell_value(r, c)
        nrows, ncols, first = sheet.nrows, sheet.ncols, 0
    else:
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(data))["기본정보"]
        get = lambda r, c: sheet.cell(r, c).value
        nrows, ncols, first = sheet.max_row + 1, sheet.max_column + 1, 1
    for header in range(first, first + 10):
        cols = find_columns([get(header, c) or "" for c in range(first, ncols)], offset=first)
        if cols:
            break
    else:
        raise SystemExit("기본정보 시트에서 상품명 열을 찾지 못했습니다.")
    val = lambda r, key: (get(r, cols[key]) or "") if cols[key] is not None else ""
    rows = []
    # 이셀러스 기본정보 시트는 제목 다음 줄이 입력 안내라서 그다음 줄부터 상품입니다.
    for r in range(header + 2, nrows):
        name = str(val(r, "name")).strip()
        if name:
            rows.append({"row": r, "excel_row": r + (1 if first == 0 else 0), "no": code_text(val(r, "no")),
                         "old": name, "price": val(r, "price"),
                         "quantity": val(r, "quantity"), "origin": str(val(r, "origin")).strip(),
                         "folder": str(val(r, "folder")).strip(), "option_details": str(val(r, "option_details"))})
    return rows, cols


def code_text(value) -> str:
    """엑셀이 숫자로 읽은 번호(68033024.0)를 글자('68033024')로."""
    text = str(value or "").strip()
    return text[:-2] if text.endswith(".0") else text


def ext_rows(path: Path) -> dict[str, int]:
    """확장정보 시트: 원본번호 → 행 위치(xls는 0부터, xlsx는 1부터). 시트가 없으면 {}."""
    data = path.read_bytes()
    if path.suffix.lower() == ".xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        if "확장정보" not in book.sheet_names():
            return {}
        sheet = book.sheet_by_name("확장정보")
        return {code_text(sheet.cell_value(r, 0)): r for r in range(2, sheet.nrows) if code_text(sheet.cell_value(r, 0))}
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data))
    if "확장정보" not in book.sheetnames:
        return {}
    sheet = book["확장정보"]
    return {code_text(sheet.cell(r, 1).value): r for r in range(3, sheet.max_row + 1) if code_text(sheet.cell(r, 1).value)}


def ext_market_text(path: Path) -> dict[str, str]:
    """확장정보 원본번호 → '마켓 카테고리번호'(B열) 글자. 저장 후 확인용."""
    data, rows = path.read_bytes(), ext_rows(path)
    if path.suffix.lower() == ".xls":
        import xlrd
        sheet = xlrd.open_workbook(file_contents=data).sheet_by_name("확장정보") if rows else None
        return {no: str(sheet.cell_value(r, 1)) for no, r in rows.items()}
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(data))["확장정보"] if rows else None
    return {no: str(sheet.cell(r, 2).value or "") for no, r in rows.items()}


def number(value) -> Decimal | None:
    try:
        return Decimal(str(value).replace(",", "").strip()) if str(value).strip() else None
    except InvalidOperation:
        return None


DEFAULT_PRICING = {"markup": 1.25, "fee_rate": 0.06, "min_margin": 1000, "max_bundle": 5}


def pricing() -> dict:
    """판매가 규칙: 원가×markup(10원 단위), 수수료 fee_rate를 떼고 한 건당 min_margin원 이상 남게."""
    return {**DEFAULT_PRICING, **config().get("pricing", {})}


def margin(price, cost) -> Decimal:
    """한 건 팔았을 때 남는 돈 = 판매가 - 수수료 - 원가 (택배비·포장비·부가세는 빼지 않음)."""
    return Decimal(price) * (1 - Decimal(str(pricing()["fee_rate"]))) - Decimal(cost)


def price_plan(value) -> dict | None:
    """원가(엑셀 판매가) → {"bundle": 묶음 개수, "price": 새 판매가, "margin": 남는 돈}.
    원가×25%로 마진이 모자라면 묶음 개수를 늘리고(max_bundle까지), 그래도 모자라면 판매가를 올립니다."""
    cost = number(value)
    if cost is None:
        return None
    rule = pricing()
    markup, keep = Decimal(str(rule["markup"])), Decimal(str(rule["min_margin"]))
    fee = Decimal(str(rule["fee_rate"]))
    for n in range(1, max(1, int(rule["max_bundle"])) + 1):
        price = (cost * n * markup).quantize(Decimal("1E1"), rounding=ROUND_HALF_UP)
        if margin(price, cost * n) >= keep:
            break
    else:
        # 최대 묶음으로도 모자라면 마진을 맞추는 가장 낮은 10원 단위 판매가로 올립니다.
        price = Decimal(math.ceil((cost * n + keep) / (1 - fee) / 10) * 10)
    return {"bundle": n, "price": int(price), "margin": int(margin(price, cost * n))}


def bundle_suffix(n: int) -> str:
    """상품명 끝에 붙이는 묶음 개수(1개면 붙이지 않음)."""
    return f" {n}개" if n > 1 else ""


def option_stock(value) -> int | None:
    total, found = 0, False
    for line in str(value or "").replace("\r", "\n").split("\n"):
        parts = line.strip().split("*")
        if len(parts) >= 4 and number(parts[3]) is not None:
            total += int(number(parts[3]))
            found = True
    return total if found else None


def updates_for(row: dict, cols: dict, title: str | None, rules: bool = True) -> dict:
    """title에는 묶음 개수가 이미 붙어 있어야 합니다. 수량은 묶음 단위(재고 ÷ 묶음 개수)로 적습니다."""
    cells = {}
    if title:
        cells[cols["name"]] = title
    if rules:
        plan = price_plan(row["price"])
        if plan is not None and cols["price"] is not None:
            cells[cols["price"]] = plan["price"]
        bundle = plan["bundle"] if plan else 1
        stock = option_stock(row["option_details"])
        qty = stock if stock is not None else number(row["quantity"])
        if qty is not None and cols["quantity"] is not None:
            cells[cols["quantity"]] = int(qty) // bundle
        if row["origin"] == "기타" and cols["origin"] is not None:
            cells[cols["origin"]] = "수입/아시아/중국"
    return cells


def write_excel(src: Path, dst: Path, updates: dict[int, dict], market_text: dict[str, str] | None = None) -> None:
    """updates = 기본정보 {행: {열: 값}}, market_text = 확장정보 {원본번호: B열 '마켓명*번호' 줄들}."""
    data = src.read_bytes()
    ext = ext_rows(src) if market_text else {}
    if src.suffix.lower() == ".xls":
        import xlrd
        from xlutils.copy import copy
        book = xlrd.open_workbook(file_contents=data, formatting_info=True)
        out = copy(book)
        sheet = out.get_sheet(book.sheet_names().index("기본정보"))
        for r, cells in updates.items():
            for c, v in cells.items():
                sheet.write(r, c, v)
        if ext:
            ext_sheet = out.get_sheet(book.sheet_names().index("확장정보"))
            for no, text in market_text.items():
                ext_sheet.write(ext[no], 1, text)
        out.save(str(dst))
    else:
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data))
        sheet = book["기본정보"]
        for r, cells in updates.items():
            for c, v in cells.items():
                sheet.cell(r, c).value = v
        for no, text in market_text.items() if ext else ():
            book["확장정보"].cell(ext[no], 2).value = text
        book.save(dst)


# ---------------------------------------------------------------- Itemscout keyword DB

class KeywordDB:
    def __init__(self, path: Path = DB_PATH):
        self.rows, self.exact, self.index = [], {}, defaultdict(list)
        with gzip.open(path, "rt", encoding="utf-8", newline="") as handle:
            for rec in csv.DictReader(handle):
                key = compact(rec["keyword"])
                if not key or key in self.exact:
                    continue
                idx = len(self.rows)
                self.rows.append((rec["keyword"], key, rec["category"], int(rec["volume"] or 0),
                                  int(rec["products"]) if rec["products"] else None,
                                  float(rec["competition"]) if rec["competition"] else None))
                self.exact[key] = idx
                for gram in {key[i:i + 2] for i in range(len(key) - 1)}:
                    self.index[gram].append(idx)

    @staticmethod
    def score(row) -> float:
        comp = 5.0 if row[5] is None else row[5]
        return math.log10(row[3] + 10) - 0.6 * math.log10(comp + 0.1)

    def get(self, term):
        idx = self.exact.get(compact(term))
        return self.rows[idx] if idx is not None else None

    def containing(self, term, limit=12):
        key = compact(term)
        if len(key) < 2:
            return []
        postings = min((self.index.get(key[i:i + 2], []) for i in range(len(key) - 1)), key=len)
        found = [self.rows[i] for i in postings if key in self.rows[i][1]]
        return sorted(found, key=self.score, reverse=True)[:limit]


def fmt(row) -> str:
    comp = "-" if row[5] is None else f"{row[5]:g}"
    cat = " > ".join(row[2].split(" > ")[-2:]) if row[2] else "-"
    return f"{row[0]} 검색{row[3]:,}/경쟁{comp} [{cat}]"


# ---------------------------------------------------------------- title checks

def longest_common(a: str, b: str) -> str:
    best = ""
    for i in range(len(a)):
        for j in range(i + len(best) + 1, len(a) + 1):
            if a[i:j] in b:
                best = a[i:j]
            else:
                break
    return best


_db = None


def is_word(part: str) -> bool:
    """아이템스카우트 DB에 키워드로 있으면 실제 단어로 봅니다(DB가 없으면 단어로 간주)."""
    global _db
    if _db is None:
        _db = KeywordDB() if DB_PATH.exists() else False
    return True if _db is False else _db.get(part) is not None


def check_title(title: str, source: str, suffix: str = "") -> tuple[list[str], list[str]]:
    """(오류, 참고) 반환. 오류가 있으면 저장하지 않습니다. suffix = 뒤에 붙일 묶음 개수(' 3개')."""
    errors, notes = [], []
    words = title.split()
    src_words = clean(source).split()
    src_key = compact(source)
    if len(title + suffix) > TITLE_MAX:
        errors.append(f"{len(title + suffix)}자 (40자 초과" + (f", 묶음 '{suffix.strip()}' 포함)" if suffix else ")"))
    pos = 0
    for w in src_words:
        try:
            pos = words.index(w, pos) + 1
        except ValueError:
            errors.append(f"기존 단어 빠짐/순서 바뀜: {w}")
    for i, a in enumerate(words):
        ka = compact(a)
        if len(ka) < 2 and a not in src_words:
            errors.append(f"한 글자 단어: {a}")
        for b in words[i + 1:]:
            kb = compact(b)
            if ka == kb:
                errors.append(f"같은 단어 반복: {a}")
            elif a in src_words and b in src_words:
                continue  # 사장님이 정리한 기존 상품명끼리는 검사하지 않음
            elif len(ka) >= 2 and len(kb) >= 2 and (ka in kb or kb in ka):
                errors.append(f"단어 포함 중복: {a} / {b}")
            elif len(longest_common(ka, kb)) >= 3 and is_word(longest_common(ka, kb)):
                # 겹친 부분이 실제 단어일 때만(텀블러가방/텀블러홀더 → '텀블러') 반복으로 봅니다.
                # 전동드라이버/전동드릴의 '전동드'처럼 우연히 겹친 글자는 괜찮습니다.
                errors.append(f"같은 단어가 두 번 들어감: {a} / {b} ('{longest_common(ka, kb)}')")
        if a not in src_words:
            if any(b.lower() in ka for b in map(compact, brands()) if b):
                errors.append(f"브랜드: {a}")
            if any(p in ka for p in PROMO_WORDS):
                errors.append(f"과장 문구: {a}")
            if re.search(r"\d", ka) and ka not in src_key:
                errors.append(f"원본에 없는 숫자: {a}")
            if any((ka == c or ka.endswith(c)) and c not in src_key for c in COLOR_WORDS + SIZE_WORDS):
                errors.append(f"원본에 없는 색상/크기: {a}")
    if len(title + suffix) < TITLE_MIN:
        notes.append(f"{len(title + suffix)}자 · 30자 미만(쓸 단어 부족) → 직접 보충 필요")
    return errors, notes
